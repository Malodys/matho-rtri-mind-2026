# -*- coding: utf-8 -*-
# AI 辅助：Codex (GPT-6), OpenAI, 2026-10-30
# 人工复核状态：待团队核验；以下为产物一致性自动检查。
import unittest
import subprocess
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"data/processed/q1_tabpfn_final"

class Q1PaperRevisionTests(unittest.TestCase):
    # AI 辅助实现，尚未人工复核。
    def test_metrics_match_archived_final_experiment(self):
        row=pd.read_csv(OUT/"model_summary.csv").set_index("model").loc["tabpfn"]
        np.testing.assert_allclose([row.R2_mean,row.Macro_MAE,row.Macro_RMSE],
            [0.2407293912437345,0.3362367581706291,0.47470748219473685],atol=1e-12)
    def test_holdout_labels_and_sample_ids(self):
        p=pd.read_csv(OUT/"test_predictions.csv.gz")
        y=pd.read_csv(ROOT/"data/splits/seed_0/4_test.csv").set_index("beer")
        self.assertEqual(set(p.beer),set(y.index))
        self.assertEqual(len(p),3750)
        self.assertFalse(p.duplicated(["beer","target"]).any())
        np.testing.assert_allclose(p.actual,[y.at[b,t] for b,t in zip(p.beer,p.target)],atol=1e-14)
    def test_grade_cutpoints_match_original_protocol(self):
        y=pd.read_csv(ROOT/"data/splits/seed_0/4_train.csv")
        cuts=pd.read_csv(ROOT/"data/processed/q1_final/proxy_grade_cutpoints.csv",index_col=0)
        for t in cuts.index:
            np.testing.assert_allclose(np.quantile(y[t],np.arange(1,7)/7),cuts.loc[t],atol=1e-12)
        p=pd.read_csv(OUT/"sensory_scores_0_6.csv")
        grades=p[[c for c in p if c.endswith("_grade")]]
        self.assertEqual(grades.shape,(250,50))
        self.assertTrue(grades.notna().all().all())
        self.assertTrue(((grades>=0)&(grades<=6)).all().all())
    def test_collaborator_original_is_unchanged(self):
        for name in ["README.md","paper.docx","paper.pdf","paper_check.json","paper_check.txt"]:
            rel="data/processed/paper/"+name
            remote=subprocess.check_output(["git","show","773269f:"+rel],cwd=ROOT)
            actual=(ROOT/rel).read_bytes()
            if name.endswith((".md",".json",".txt")):
                actual=actual.replace(b"\r\n",b"\n")
                remote=remote.replace(b"\r\n",b"\n")
            self.assertEqual(actual,remote,name)
    def test_collaborator_validation_is_in_revised_paper(self):
        d=Document(ROOT/"data/processed/paper-q1-revised/paper.docx")
        text="".join(p.text for p in d.paragraphs).replace(" ","")
        for value in ["83/150","30/32","FDR","48","12.7","0.240729"]:
            self.assertIn(value,text)
    def test_revised_prose_has_no_redundant_exclusion_claims(self):
        d=Document(ROOT/"data/processed/paper-q1-revised/paper.docx")
        text="".join(p.text for p in d.paragraphs)
        for phrase in ["不声称","不采用条件生成对抗网络","不验证化学机理","不将其解释为","不作市场因果解释","需要强调其边界","，。"]:
            self.assertNotIn(phrase,text)
        for value in ["83/150","30/32","42/48","0.240729","13.5%"]:
            self.assertIn(value,text.replace(" ",""))
    def test_body_uses_academic_method_descriptions(self):
        d=Document(ROOT/"data/processed/paper-q1-revised/paper.docx")
        body=[]
        for p in d.paragraphs:
            if p.text.strip()=="参考文献":break
            body.append(p.text)
        text="".join(body)
        for phrase in ["恢复作者","沿用","原论文","原稿","原方案","队友","与原文","叙事主线","本文全部引用按","模型选择分两步进行"]:
            self.assertNotIn(phrase,text)
        self.assertIn("研究背景",text)
        self.assertIn("共享参数GBR","".join(c.text for t in d.tables for r in t.rows for c in r.cells))
    def test_document_uses_final_method_and_metrics(self):
        d=Document(ROOT/"data/processed/paper-q1-revised/paper.docx")
        text="\n".join(p.text for p in d.paragraphs)
        self.assertIn("0.240729",text)
        self.assertNotIn("TabPFN重复",text.replace(" ",""))
        summary=pd.read_csv(OUT/"combined_summary.csv")
        self.assertEqual(int(summary.model.eq("tabpfn").sum()),1)
        self.assertNotIn("tabpfn_uniform",set(summary.model))
        matches=[row for table in d.tables for row in table.rows if row.cells[0].text.strip()=="TabPFN"]
        self.assertEqual(len(matches),1)
        self.assertEqual(matches[0].cells[1].text.strip(),"0.2407")
        q1=text.split("5.1 问题一：感官评分预测与 0–6 分制")[-1].split("5.2 消费者感知路径与市场价值")[0]
        self.assertNotIn("RMSE",q1)
        self.assertNotIn("MAE",q1)
        self.assertNotIn("Kappa",q1)
        self.assertNotIn("不将其与主模型直接比较",text)
        self.assertEqual(sum(p.text.replace(" ","")=="目录" for p in d.paragraphs),1)

if __name__=="__main__": unittest.main()
